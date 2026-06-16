# -*- coding: utf-8 -*-
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

PATH = r"C:\Users\alexa\OneDrive\Desktop\alexandre\Sistema\V3\GERAÇÃO DE MATERIAL.xlsm"

print("Loading (data_only=True, cached values)...", flush=True)
wb = openpyxl.load_workbook(PATH, data_only=True, read_only=True)

for ws in wb.worksheets:
    print("="*100)
    print(f"SHEET: {ws.title!r}   max_row={ws.max_row} max_col={ws.max_column} state={ws.sheet_state}")
    # print first 8 rows x first 14 cols of cached values
    maxr = min(ws.max_row or 1, 8)
    maxc = min(ws.max_column or 1, 14)
    for r in range(1, maxr+1):
        cells = []
        for c in range(1, maxc+1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                cells.append("")
            else:
                s = str(v).replace("\n", " ")
                if len(s) > 18: s = s[:17]+"~"
                cells.append(s)
        if any(cells):
            print(f"  R{r}: " + " | ".join(cells))
wb.close()
print("DONE")
