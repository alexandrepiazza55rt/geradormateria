# -*- coding: utf-8 -*-
import openpyxl, sys, io
from openpyxl.utils import get_column_letter
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
PATH = r"C:\Users\alexa\OneDrive\Desktop\alexandre\Sistema\V3\GERAÇÃO DE MATERIAL.xlsm"
wb = openpyxl.load_workbook(PATH, data_only=True, read_only=True)

def row_labels(ws, row, c1, c2):
    out=[]
    for c in range(c1,c2+1):
        v=ws.cell(row=row,column=c).value
        if v not in (None,""):
            out.append(f"{get_column_letter(c)}={str(v).strip()}")
    return out

# 1) Structure headers (row 3) for the 4 generators
for name in ['Mono 13,8kv','Mono 34,5kv','Trifásico 13,8kv ','Trifásico 34,5kv']:
    ws=wb[name]
    print(f"\n=== {name}: ROW3 structure blocks (cols up to {ws.max_column}) ===")
    print("  " + " | ".join(row_labels(ws,3,1,ws.max_column)))
    print(f"--- {name}: ROW4 sub-headers ---")
    print("  " + " | ".join(row_labels(ws,4,4,ws.max_column)))

# 2) Estruturas sheets: section labels in col E (structure catalog)
for name in ['Estruturas 13,8kv','Estruturas 34,5kv']:
    ws=wb[name]
    print(f"\n=== {name}: structure section labels (col E where label is a group header) ===")
    labels=[]
    for r in range(1, min(ws.max_row,2300)+1):
        e=ws.cell(row=r,column=5).value  # col E
        d=ws.cell(row=r,column=4).value  # col D
        # group headers: col E has text that's not a pole code (poles start with DT or digit/)
        if isinstance(e,str) and e.strip():
            s=e.strip()
            if not (s.startswith('DT-') or (len(s)>1 and s[0].isdigit() and '/' in s)):
                labels.append(f"R{r}:{s}")
    print("  " + " | ".join(labels))

# 3) Materiais e Valores: count materials, detect blocks
ws=wb['Materiais e Valores']
mats=0; first=None; last=None
for r in range(1, ws.max_row+1):
    c=ws.cell(row=r,column=3).value  # descrição col C
    if isinstance(c,str) and c.strip() and r>=19:
        mats+=1
        if first is None: first=r
        last=r
print(f"\n=== Materiais e Valores: {mats} material rows with description, between R{first} and R{last} (max_row={ws.max_row}) ===")
# print col headers row 18
print("  Header R18: " + " | ".join(row_labels(ws,18,1,13)))

# 4) Estruturas dims
print(f"\nEstruturas 13,8kv max_col={wb['Estruturas 13,8kv'].max_column} max_row={wb['Estruturas 13,8kv'].max_row}")
wb.close()
print("DONE")
