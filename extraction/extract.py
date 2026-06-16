# -*- coding: utf-8 -*-
"""
Fase 2 extraction driver. Stages:
  master   -> material master list + output map (which sheet/col/row feeds each material)
  (more stages added incrementally)

The output map is the rigorous backbone: we parse every formula in
'Materiais e Valores'!E19:E1255 to learn, for each master material, the exact
list of (sheet, col, row) cells it sums. This tells us each contributing
sheet's "total column" and the sheet-row -> material mapping, with no guessing.
"""
import sys, io, json, os, time
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlmodel import Model, tokenize, Parser
from openpyxl.utils import coordinate_to_tuple, get_column_letter

PATH = r"C:\Users\alexa\OneDrive\Desktop\alexandre\Sistema\V3\GERAÇÃO DE MATERIAL.xlsm"
OUTDIR = r"C:\Users\alexa\OneDrive\Desktop\alexandre\Sistema\V3\data"
MV = "Materiais e Valores"
os.makedirs(OUTDIR, exist_ok=True)

def collect_refs(node, out):
    t = node[0]
    if t == "cell": out.append(("cell", node[1], node[2]))
    elif t == "range": out.append(("range", node[1], node[2], node[3]))
    elif t in ("bin", "cmp", "concat"): collect_refs(node[2], out); collect_refs(node[3], out)
    elif t in ("unary", "percent"): collect_refs(node[-1], out)
    elif t == "func":
        for a in node[2]: collect_refs(a, out)

def build_master(m):
    """Return list of materials (master) and the output map."""
    materials = []
    out_map = {}   # master_id -> list of (sheet, coord)
    # find the data rows: from 19 down to where there is a description in col C
    for r in range(19, 1256):
        desc = m.literal.get((MV, f"C{r}"))
        sapv = m.literal.get((MV, f"B{r}"))
        lider = m.literal.get((MV, f"A{r}"))
        unit = m.literal.get((MV, f"D{r}"))
        has_formula = (MV, f"E{r}") in m.formula
        if desc is None and not has_formula:
            continue
        mid = len(materials) + 1
        materials.append({
            "id": mid,
            "row": r,
            "cod_sap": ("" if sapv is None else str(sapv).strip()),
            "cod_lider7": ("" if lider is None else str(lider).strip()),
            "descricao": ("" if desc is None else str(desc)).strip(),
            "unidade": ("" if unit is None else str(unit)).strip(),
        })
        # parse E formula refs
        refs = []
        f = m.formula.get((MV, f"E{r}"))
        if f and "[" not in f:
            try:
                ast = Parser(tokenize(f), MV).parse()
                rr = []
                collect_refs(ast, rr)
                for ref in rr:
                    if ref[0] == "cell":
                        refs.append((ref[1], ref[2]))
                    elif ref[0] == "range":
                        # expand small ranges
                        (r1, c1) = coordinate_to_tuple(ref[2]); (r2, c2) = coordinate_to_tuple(ref[3])
                        for cc in range(min(c1,c2), max(c1,c2)+1):
                            for rw in range(min(r1,r2), max(r1,r2)+1):
                                refs.append((ref[1], f"{get_column_letter(cc)}{rw}"))
            except Exception as e:
                refs = [("PARSE_ERROR", str(e))]
        out_map[mid] = refs
    return materials, out_map

def main():
    t0 = time.time()
    print("loading model...", flush=True)
    m = Model(PATH)
    print(f"loaded {time.time()-t0:.1f}s", flush=True)
    materials, out_map = build_master(m)
    print(f"materials: {len(materials)}")

    # analyze: which sheets/columns feed the output, and row mapping per sheet
    from collections import defaultdict, Counter
    sheet_cols = defaultdict(Counter)        # sheet -> Counter(col)
    sheet_rowmap = defaultdict(dict)         # sheet -> {sheet_row: set(master_id)}
    nrefs = Counter()
    for mat in materials:
        refs = out_map[mat["id"]]
        nrefs[len(refs)] += 1
        for (sheet, coord) in refs:
            if sheet == "PARSE_ERROR":
                sheet_cols["PARSE_ERROR"][coord[:40]] += 1
                continue
            (row, col) = coordinate_to_tuple(coord)
            sheet_cols[sheet][get_column_letter(col)] += 1
            sheet_rowmap[sheet].setdefault(row, set()).add(mat["id"])

    print("\n=== Sheets feeding the consolidated BOM (col usage counts) ===")
    for sheet in sorted(sheet_cols):
        cols = sheet_cols[sheet]
        print(f"  {sheet:32} -> {dict(cols)}")
    print("\n=== refs-per-material distribution ===", dict(nrefs))

    # save
    with io.open(os.path.join(OUTDIR, "materiais.json"), "w", encoding="utf-8") as fh:
        json.dump(materials, fh, ensure_ascii=False, indent=1)
    # output map: master_id -> [[sheet,coord],...]
    om = {str(mid): [[s, c] for (s, c) in refs] for mid, refs in out_map.items()}
    with io.open(os.path.join(OUTDIR, "output_map.json"), "w", encoding="utf-8") as fh:
        json.dump(om, fh, ensure_ascii=False)
    print(f"\nwrote materiais.json ({len(materials)}), output_map.json. total {time.time()-t0:.1f}s")

if __name__ == "__main__":
    main()
