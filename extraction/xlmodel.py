# -*- coding: utf-8 -*-
"""
xlmodel - a small, dependency-light evaluator for GERAÇÃO DE MATERIAL.xlsm.

The workbook is (almost entirely) a linear system whose only functions are
SUM, IF, AND, OR, ROUNDUP, ROUNDDOWN plus + - * / and comparisons.  This module
loads the workbook once (formulas + cached values via openpyxl), parses each
formula into a tiny AST, and evaluates any cell with support for input
*overrides* (so we can probe the model: set inputs, read outputs).

Fidelity rules:
  * empty cell -> 0 in numeric context (matches Excel for +,-,*,SUM).
  * a formula that references an EXTERNAL workbook ([1]..., [2]...) cannot be
    recomputed (we don't have the external file) -> we return its cached value.
    These are only 11 cells, all in 'Transformadores'.
"""
import re, math
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

# ----------------------------- tokenizer ------------------------------------
TOKEN_RE = re.compile(r"""
    (?P<WS>\s+)
  | (?P<NUM>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)
  | (?P<STR>"(?:[^"]|"")*")
  | (?P<SHEETQ>'(?:[^']|'')*'!)            # 'Sheet Name'!
  | (?P<EXT>\[\d+\])                       # external link marker [1]
  | (?P<CELL>\$?[A-Za-z]{1,3}\$?\d+)       # cell ref incl. absolute $A$1
  | (?P<NAME>[A-Za-z_][A-Za-z0-9_\.]*)
  | (?P<OP><=|>=|<>|[-+*/^=<>():;,!&%])
""", re.VERBOSE)

def tokenize(s):
    toks = []
    i = 0
    while i < len(s):
        m = TOKEN_RE.match(s, i)
        if not m:
            raise SyntaxError(f"bad token at {i!r} in {s!r}")
        i = m.end()
        kind = m.lastgroup
        if kind == "WS":
            continue
        toks.append((kind, m.group()))
    toks.append(("EOF", ""))
    return toks

# ----------------------------- parser ---------------------------------------
# Grammar (precedence low->high):
#   expr     := cmp
#   cmp      := concat (('='|'<>'|'<='|'>='|'<'|'>') concat)*
#   concat   := add ('&' add)*
#   add      := mul (('+'|'-') mul)*
#   mul      := unary (('*'|'/') unary)*
#   unary    := ('-'|'+')? power
#   power    := primary ('^' unary)?
#   primary  := NUM | STR | '(' expr ')' | func | ref
#   ref      := [EXT] [SHEET] CELL [':' [SHEET] CELL]
#   func     := NAME '(' [expr (','|';' expr)*] ')'
CELL_RE = re.compile(r"^\$?([A-Z]{1,3})\$?(\d+)$")

class Parser:
    def __init__(self, toks, cur_sheet):
        self.toks = toks
        self.p = 0
        self.cur_sheet = cur_sheet

    def peek(self): return self.toks[self.p]
    def nextt(self):
        t = self.toks[self.p]; self.p += 1; return t
    def expect(self, val):
        k, v = self.toks[self.p]
        if v != val: raise SyntaxError(f"expected {val!r} got {v!r}")
        self.p += 1

    def parse(self):
        node = self.expr()
        if self.peek()[0] != "EOF":
            raise SyntaxError(f"trailing tokens at {self.p}: {self.toks[self.p:]}")
        return node

    def expr(self): return self.cmp()

    def cmp(self):
        node = self.concat()
        while self.peek()[0] == "OP" and self.peek()[1] in ("=", "<>", "<=", ">=", "<", ">"):
            op = self.nextt()[1]
            rhs = self.concat()
            node = ("cmp", op, node, rhs)
        return node

    def concat(self):
        node = self.add()
        while self.peek()[1] == "&":
            self.nextt(); rhs = self.add(); node = ("concat", node, rhs)
        return node

    def add(self):
        node = self.mul()
        while self.peek()[1] in ("+", "-"):
            op = self.nextt()[1]; rhs = self.mul(); node = ("bin", op, node, rhs)
        return node

    def mul(self):
        node = self.unary()
        while self.peek()[1] in ("*", "/"):
            op = self.nextt()[1]; rhs = self.unary(); node = ("bin", op, node, rhs)
        return node

    def unary(self):
        if self.peek()[1] in ("-", "+"):
            op = self.nextt()[1]; node = self.unary(); return ("unary", op, node)
        return self.power()

    def power(self):
        node = self.primary()
        while self.peek()[1] == "%":
            self.nextt(); node = ("percent", node)
        if self.peek()[1] == "^":
            self.nextt(); rhs = self.unary(); node = ("bin", "^", node, rhs)
        return node

    def primary(self):
        k, v = self.peek()
        if k == "NUM":
            self.nextt(); return ("num", float(v))
        if k == "STR":
            self.nextt(); return ("str", v[1:-1].replace('""', '"'))
        if v == "(":
            self.nextt(); node = self.expr(); self.expect(")"); return node
        if k == "EXT":
            # external reference -> whole sub-expression unusable; mark
            self.nextt()
            # consume an optional sheet + cell/range following it
            self._consume_ref_tail()
            return ("extern",)
        if k == "SHEETQ":
            sheet = v[1:-2].replace("''", "'")  # strip quotes and trailing !
            self.nextt()
            return self._ref_after_sheet(sheet)
        if k == "CELL":
            self.nextt()
            return self._make_ref_or_range(self.cur_sheet, v)
        if k == "NAME":
            # could be a function call, or a sheet name (Sheet!Cell)
            name = v
            self.nextt()
            if self.peek()[1] == "(":
                return self._func(name)
            if self.peek()[1] == "!":
                self.nextt()
                return self._ref_after_sheet(name.strip())
            raise SyntaxError(f"unexpected NAME {name!r} (not func/sheet/cell)")
        raise SyntaxError(f"unexpected token {k}:{v!r}")

    def _consume_ref_tail(self):
        # after [n], expect optional SHEET then CELL[:CELL]; consume tokens
        if self.peek()[0] == "SHEETQ":
            self.nextt()
        elif self.peek()[0] == "NAME" and self.toks[self.p+1][1] == "!":
            self.nextt(); self.nextt()
        if self.peek()[0] in ("CELL", "NAME"):
            self.nextt()
            if self.peek()[1] == ":":
                self.nextt()
                if self.peek()[0] == "EXT": self.nextt()
                if self.peek()[0] == "SHEETQ": self.nextt()
                if self.peek()[0] in ("CELL", "NAME"): self.nextt()

    def _ref_after_sheet(self, sheet):
        k, v = self.peek()
        if k not in ("CELL", "NAME"):
            raise SyntaxError(f"expected cell after sheet, got {v!r}")
        self.nextt()
        return self._make_ref_or_range(sheet, v)

    def _make_ref_or_range(self, sheet, first):
        first = first.replace("$", "").upper()
        if self.peek()[1] == ":":
            self.nextt()
            k2, v2 = self.peek()
            if k2 == "SHEETQ":
                self.nextt(); k2, v2 = self.peek()
            elif k2 == "NAME" and self.toks[self.p+1][1] == "!":
                self.nextt(); self.nextt(); k2, v2 = self.peek()
            self.nextt()  # consume second cell
            return ("range", sheet, first, v2.replace("$", "").upper())
        return ("cell", sheet, first)

    def _arg(self):
        # tolerate empty arguments, e.g. SUM(,A1:A9) or IF(a,,b)
        if self.peek()[1] in (",", ";", ")"):
            return ("num", 0.0)
        return self.expr()

    def _func(self, name):
        self.expect("(")
        args = []
        if self.peek()[1] != ")":
            args.append(self._arg())
            while self.peek()[1] in (",", ";"):
                self.nextt(); args.append(self._arg())
        self.expect(")")
        return ("func", name.upper().strip(), args)

# ----------------------------- model ----------------------------------------
class Model:
    def __init__(self, path):
        fwb = openpyxl.load_workbook(path, data_only=False)
        vwb = openpyxl.load_workbook(path, data_only=True)
        self.formula = {}   # (sheet,coord) -> raw formula string (without '=') or None
        self.literal = {}   # (sheet,coord) -> literal value (number/str)
        self.cached  = {}   # (sheet,coord) -> cached value
        self.ast_cache = {}
        self.sheets = [ws.title for ws in fwb.worksheets]
        for ws in fwb.worksheets:
            s = ws.title
            for row in ws.iter_rows():
                for c in row:
                    val = c.value
                    if val is None:
                        continue
                    key = (s, c.coordinate)
                    if isinstance(val, str) and val.startswith("="):
                        self.formula[key] = val[1:]
                    else:
                        self.literal[key] = val
        for ws in vwb.worksheets:
            s = ws.title
            for row in ws.iter_rows():
                for c in row:
                    if c.value is not None:
                        self.cached[(s, c.coordinate)] = c.value
        self.overrides = {}
        self.cache = {}

    # ---- input control ----
    def set_inputs(self, mapping):
        """mapping: {(sheet,coord): value}. Resets evaluation cache."""
        self.overrides = dict(mapping)
        self.cache = {}

    def clear(self):
        self.overrides = {}; self.cache = {}

    # ---- evaluation ----
    def get_ast(self, sheet, formula):
        key = (sheet, formula)
        a = self.ast_cache.get(key)
        if a is None:
            a = Parser(tokenize(formula), sheet).parse()
            self.ast_cache[key] = a
        return a

    def cell(self, sheet, coord):
        coord = coord.upper()
        key = (sheet, coord)
        if key in self.overrides:
            return self.overrides[key]
        if key in self.cache:
            return self.cache[key]
        # guard against accidental cycles
        self.cache[key] = 0.0
        f = self.formula.get(key)
        if f is None:
            v = self.literal.get(key, 0.0)
        else:
            if "[" in f:                     # external reference -> use cached value
                v = self.cached.get(key, 0.0) or 0.0
            else:
                try:
                    v = self.eval(self.get_ast(sheet, f), sheet)
                except Exception as e:
                    raise RuntimeError(f"eval failed at {sheet}!{coord}: {f}\n  -> {e}")
        self.cache[key] = v
        return v

    def _num(self, v):
        if v is None or v == "":
            return 0.0
        if isinstance(v, bool):
            return 1.0 if v else 0.0
        if isinstance(v, (int, float)):
            return float(v)
        return 0.0   # text in numeric context -> 0 (we never sum header text)

    def _iter_range(self, sheet, a, b):
        m1 = CELL_RE.match(a); m2 = CELL_RE.match(b)
        c1 = column_index_from_string(m1.group(1)); r1 = int(m1.group(2))
        c2 = column_index_from_string(m2.group(1)); r2 = int(m2.group(2))
        if c1 > c2: c1, c2 = c2, c1
        if r1 > r2: r1, r2 = r2, r1
        for cc in range(c1, c2 + 1):
            col = get_column_letter(cc)
            for rr in range(r1, r2 + 1):
                yield (sheet, f"{col}{rr}")

    def eval(self, node, sheet):
        t = node[0]
        if t == "num": return node[1]
        if t == "str": return node[1]
        if t == "extern":
            return 0.0  # handled at cell level; bare extern inside expr -> 0
        if t == "cell":
            return self.cell(node[1], node[2])
        if t == "range":
            # bare range used as scalar -> sum (defensive); normally only inside SUM
            return sum(self._num(self.cell(s, c)) for (s, c) in self._iter_range(node[1], node[2], node[3]))
        if t == "percent":
            return self._num(self.eval(node[1], sheet)) / 100.0
        if t == "unary":
            v = self._num(self.eval(node[2], sheet))
            return -v if node[1] == "-" else v
        if t == "bin":
            op = node[1]
            a = self._num(self.eval(node[2], sheet)); b = self._num(self.eval(node[3], sheet))
            if op == "+": return a + b
            if op == "-": return a - b
            if op == "*": return a * b
            if op == "/": return a / b if b != 0 else 0.0
            if op == "^": return a ** b
        if t == "concat":
            return str(self.eval(node[1], sheet)) + str(self.eval(node[2], sheet))
        if t == "cmp":
            a = self.eval(node[2], sheet); b = self.eval(node[3], sheet)
            an = self._num(a) if not isinstance(a, str) else a
            bn = self._num(b) if not isinstance(b, str) else b
            op = node[1]
            try:
                if op == "=":  return an == bn
                if op == "<>": return an != bn
                if op == "<":  return an < bn
                if op == ">":  return an > bn
                if op == "<=": return an <= bn
                if op == ">=": return an >= bn
            except TypeError:
                return False
        if t == "func":
            return self._func(node[1], node[2], sheet)
        raise RuntimeError(f"unknown node {node!r}")

    def _func(self, name, args, sheet):
        if name == "SUM":
            total = 0.0
            for a in args:
                if a[0] == "range":
                    for (s, c) in self._iter_range(a[1], a[2], a[3]):
                        total += self._num(self.cell(s, c))
                else:
                    total += self._num(self.eval(a, sheet))
            return total
        if name == "IF":
            cond = self.eval(args[0], sheet)
            cond = bool(cond) if not isinstance(cond, (int, float)) else cond != 0
            if cond:
                return self.eval(args[1], sheet) if len(args) > 1 else True
            else:
                return self.eval(args[2], sheet) if len(args) > 2 else False
        if name == "AND":
            return all(self._truth(self.eval(a, sheet)) for a in args)
        if name == "OR":
            return any(self._truth(self.eval(a, sheet)) for a in args)
        if name == "ROUNDUP":
            x = self._num(self.eval(args[0], sheet)); d = int(self._num(self.eval(args[1], sheet))) if len(args) > 1 else 0
            f = 10 ** d
            return math.ceil(abs(x) * f) / f * (1 if x >= 0 else -1)
        if name == "ROUNDDOWN":
            x = self._num(self.eval(args[0], sheet)); d = int(self._num(self.eval(args[1], sheet))) if len(args) > 1 else 0
            f = 10 ** d
            return math.floor(abs(x) * f) / f * (1 if x >= 0 else -1)
        if name == "ROUND":
            x = self._num(self.eval(args[0], sheet)); d = int(self._num(self.eval(args[1], sheet))) if len(args) > 1 else 0
            return round(x, d)
        raise RuntimeError(f"unsupported function {name}")

    def _truth(self, v):
        if isinstance(v, bool): return v
        if isinstance(v, (int, float)): return v != 0
        return bool(v)
