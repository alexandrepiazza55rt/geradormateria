# -*- coding: utf-8 -*-
"""
Extract the 4 MT generator sheets into structures with per-pole BOMs.
Output-map driven: for each sheet we read ALL cells it contributes to the
consolidated BOM (from output_map.json) and key results directly by master
material id. Each structure column (estrutura x condutor) is probed at every
pole row; results are factored into base (at reference pole) + per-pole deltas.

Output: data/estruturas_mt.json  and  data/insumos_mt.json (scalar inputs)
"""
import sys, io, os, json, time
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from xlmodel import Model, tokenize, Parser
from openpyxl.utils import coordinate_to_tuple, get_column_letter

PATH = r"C:\Users\alexa\OneDrive\Desktop\alexandre\Sistema\V3\GERAÇÃO DE MATERIAL.xlsm"
OUTDIR = r"C:\Users\alexa\OneDrive\Desktop\alexandre\Sistema\V3\data"

GEN = {
    "Mono 13,8kv":      dict(tensao=13.8, nominal=7.97,  fases=1),
    "Mono 34,5kv":      dict(tensao=34.5, nominal=19.92, fases=1),
    "Trifásico 13,8kv ":dict(tensao=13.8, nominal=13.8,  fases=3),
    "Trifásico 34,5kv": dict(tensao=34.5, nominal=34.5,  fases=3),
}

def collect_refs(node, out):
    t = node[0]
    if t in ("cell", "range"): out.append(node)
    elif t in ("bin","cmp","concat"): collect_refs(node[2],out); collect_refs(node[3],out)
    elif t in ("unary","percent"): collect_refs(node[-1],out)
    elif t=="func":
        for a in node[2]: collect_refs(a,out)

def load_inverse_map():
    """sheet -> {coord -> master_id} from output_map.json"""
    om = json.load(open(os.path.join(OUTDIR, "output_map.json"), encoding="utf-8"))
    inv = {}
    for mid, refs in om.items():
        for ref in refs:
            sheet, coord = ref[0], ref[1]
            if sheet == "PARSE_ERROR": continue
            inv.setdefault(sheet, {})[coord] = int(mid)
    return inv

def enum_input_columns(m, sheet, band_max=120):
    bycol = {}
    for (s, coord), f in m.formula.items():
        if s != sheet or "[" in f: continue
        try: ast = Parser(tokenize(f), s).parse()
        except Exception: continue
        refs = []; collect_refs(ast, refs)
        for r in refs:
            if r[0] == "range" and r[1] == sheet:
                (r1c, c1) = coordinate_to_tuple(r[2]); (r2c, c2) = coordinate_to_tuple(r[3])
                if min(r1c, r2c) > band_max: continue
                for cc in range(min(c1,c2), max(c1,c2)+1):
                    for rw in range(min(r1c,r2c), max(r1c,r2c)+1):
                        if rw > band_max: continue
                        if (sheet, f"{get_column_letter(cc)}{rw}") not in m.formula:
                            bycol.setdefault(cc, set()).add(rw)
            elif r[0] == "cell" and r[1] == sheet:
                (rw, cc) = coordinate_to_tuple(r[2])
                if rw <= band_max and (sheet, r[2]) not in m.formula:
                    bycol.setdefault(cc, set()).add(rw)
    return {c: sorted(rs) for c, rs in bycol.items()}

def label_left(m, sheet, col_idx, row):
    for cc in range(col_idx, 0, -1):
        v = m.literal.get((sheet, f"{get_column_letter(cc)}{row}"))
        if v not in (None, ""): return str(v).strip()
    return None

def pole_label_col(m, sheet, col_idx):
    for cc in range(col_idx, 0, -1):
        v = m.literal.get((sheet, f"{get_column_letter(cc)}5"))
        if isinstance(v, str) and (v.strip().startswith("DT-") or (len(v) > 1 and v[0].isdigit() and "/" in v)):
            return cc
    return None

def main():
    only = sys.argv[1] if len(sys.argv) > 1 else None
    t0 = time.time()
    print("loading model...", flush=True)
    m = Model(PATH)
    inv = load_inverse_map()
    print(f"loaded {time.time()-t0:.1f}s", flush=True)

    structures = []
    insumos = []
    sid = 0
    for sheet, cfg in GEN.items():
        if only and only not in sheet: continue
        totalcells = inv.get(sheet, {})           # coord -> master_id
        total_coords = list(totalcells.keys())
        fails = {"n": 0}
        def probe(coord_set_cell):
            m.set_inputs({(sheet, coord_set_cell): 1})
            res = {}
            for coord in total_coords:
                try:
                    v = m.cell(sheet, coord)
                except Exception as e:
                    fails["n"] += 1
                    if fails["n"] <= 3:
                        print(f"   [warn] eval fail {sheet}!{coord}: {str(e)[:80]}", flush=True)
                    continue
                if isinstance(v, (int, float)) and abs(v) > 1e-9:
                    res[totalcells[coord]] = round(v, 6)
            return res

        bycol = enum_input_columns(m, sheet)
        cols = sorted(bycol)
        print(f"\n### {sheet}: {len(cols)} input columns, {len(total_coords)} output cells", flush=True)
        nprobe = 0
        for ci in cols:
            L = get_column_letter(ci)
            rows = [r for r in bycol[ci] if r >= 5]
            struct = label_left(m, sheet, ci, 3)
            plc = pole_label_col(m, sheet, ci)
            cond = m.literal.get((sheet, f"{L}4"))
            cond = str(cond).strip() if cond not in (None, "") else None

            # Scalar/helper columns: no structure header OR no pole-label column -> treat each cell as insumo
            if struct is None or plc is None:
                for r in rows:
                    bom = probe(f"{L}{r}"); nprobe += 1
                    if not bom: continue
                    desc = label_left(m, sheet, ci, r) or f"{L}{r}"
                    insumos.append({"sheet": sheet, "cell": f"{L}{r}", "descricao": desc,
                                     "tensao_kv": cfg["tensao"], "fases": cfg["fases"], "bom": bom})
                continue

            per_pole = {}
            for r in rows:
                bom = probe(f"{L}{r}"); nprobe += 1
                if not bom: continue
                pv = m.literal.get((sheet, f"{get_column_letter(plc)}{r}"))
                plabel = str(pv).strip() if pv not in (None, "") else f"row{r}"
                per_pole[r] = (plabel, bom)
            if not per_pole: continue
            ref_r = min(per_pole)
            ref_label, ref_bom = per_pole[ref_r]
            postes = []
            for r in sorted(per_pole):
                plabel, bom = per_pole[r]
                delta = {}
                for k in set(bom) | set(ref_bom):
                    d = round(bom.get(k, 0) - ref_bom.get(k, 0), 6)
                    if abs(d) > 1e-9: delta[str(k)] = d
                postes.append({"poste": plabel, "delta": delta})
            sid += 1
            structures.append({
                "id": f"S{sid}", "sheet": sheet, "col": L,
                "tipo": struct, "condutor": cond,
                "tensao_kv": cfg["tensao"], "nominal_kv": cfg["nominal"], "fases": cfg["fases"],
                "poste_ref": ref_label,
                "base_bom": {str(k): v for k, v in ref_bom.items()},
                "postes": postes,
            })
        print(f"   probes={nprobe} structures={len(structures)} insumos={len(insumos)} t={time.time()-t0:.1f}s", flush=True)

    os.makedirs(OUTDIR, exist_ok=True)
    sfx = "" if not only else "_" + only.strip().replace(" ", "_").replace(",", "")
    json.dump(structures, open(os.path.join(OUTDIR, f"estruturas_mt{sfx}.json"), "w", encoding="utf-8"), ensure_ascii=False)
    json.dump(insumos, open(os.path.join(OUTDIR, f"insumos_mt{sfx}.json"), "w", encoding="utf-8"), ensure_ascii=False)
    print(f"\nwrote estruturas_mt{sfx}.json ({len(structures)}) and insumos_mt{sfx}.json ({len(insumos)}). total {time.time()-t0:.1f}s")

if __name__ == "__main__":
    main()
